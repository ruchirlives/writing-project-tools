"""Read and write the single-sheet assertions workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


FIELDNAMES = [
    "include", "id", "section", "precis", "assertion", "user_edit", "status", "evidence_or_check",
]


def read_assertions(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Assertions"] if "Assertions" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        header = next(values, None)
        if header is None:
            return []
        fields = [str(value or "") for value in header]
        if not all(field in fields for field in FIELDNAMES if field != "precis"):
            raise ValueError("Assertions workbook has missing columns")
        rows = []
        for values_row in values:
            row = dict(zip(fields, values_row))
            if any(value is not None and value != "" for value in values_row):
                rows.append({field: str(row.get(field) or "") for field in FIELDNAMES})
        return rows
    finally:
        workbook.close()


def write_assertions(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assertions"
    sheet.append(FIELDNAMES)
    for row in rows:
        sheet.append([str(row.get(field, "") or "") for field in FIELDNAMES])
    sheet.freeze_panes = "E2"
    widths = {"A": 12, "B": 10, "C": 22, "D": 32, "E": 80, "F": 60, "G": 18, "H": 55}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    table = Table(displayName="AssertionsList", ref=f"A1:H{max(2, sheet.max_row)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    if sheet.max_row == 1:
        sheet.append([""] * len(FIELDNAMES))
    sheet.add_table(table)
    workbook.save(path)
