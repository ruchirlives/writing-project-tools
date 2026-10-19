import unittest
from docx import Document
from writing_project_tools.create_docs import (
    is_separator_line,
    parse_table_row,
    add_runs_to_paragraph,
    add_markdown_to_document,
)


class TestCreateDocs(unittest.TestCase):
    def test_is_separator_line(self) -> None:
        self.assertTrue(is_separator_line("|---|---|"))
        self.assertTrue(is_separator_line("| :--- | ---: |"))
        self.assertTrue(is_separator_line("  | - | - |  "))
        self.assertFalse(is_separator_line("| Header 1 | Header 2 |"))
        self.assertFalse(is_separator_line(""))
        self.assertFalse(is_separator_line("some text"))

    def test_parse_table_row(self) -> None:
        self.assertEqual(parse_table_row("| A | B |"), ["A", "B"])
        self.assertEqual(parse_table_row("A | B"), ["A", "B"])
        self.assertEqual(parse_table_row("  |  spaced cell  | other  "), ["spaced cell", "other"])

    def test_add_runs_to_paragraph(self) -> None:
        doc = Document()
        p = doc.add_paragraph()
        add_runs_to_paragraph(p, "This is **Routine issues** and __bold2__ here.")
        
        runs = p.runs
        self.assertEqual(len(runs), 5)
        self.assertEqual(runs[0].text, "This is ")
        self.assertFalse(runs[0].bold)
        
        self.assertEqual(runs[1].text, "Routine issues")
        self.assertTrue(runs[1].bold)
        
        self.assertEqual(runs[2].text, " and ")
        self.assertFalse(runs[2].bold)
        
        self.assertEqual(runs[3].text, "bold2")
        self.assertTrue(runs[3].bold)

        self.assertEqual(runs[4].text, " here.")
        self.assertFalse(runs[4].bold)


    def test_add_markdown_to_document_table(self) -> None:
        doc = Document()
        md = (
            "# Heading 1\n\n"
            "| Header 1 | **Header 2** |\n"
            "|---|---|\n"
            "| Cell 1 | Cell **2** |\n"
            "| Cell 3 | Cell 4 |\n\n"
            "Normal paragraph **bold**."
        )
        add_markdown_to_document(doc, md)
        
        # Verify document has: Heading 1, Table, Paragraph
        self.assertEqual(len(doc.paragraphs), 2)  # Heading 1 and the final paragraph
        self.assertEqual(len(doc.tables), 1)
        
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)  # 1 header + 2 data rows
        self.assertEqual(len(table.columns), 2)
        
        # Header Row
        hdr_cells = table.rows[0].cells
        self.assertEqual(hdr_cells[0].paragraphs[0].text, "Header 1")
        self.assertTrue(hdr_cells[0].paragraphs[0].runs[0].bold)  # Headers are forced bold
        self.assertEqual(hdr_cells[1].paragraphs[0].text, "Header 2")
        
        # Data Rows
        self.assertEqual(table.rows[1].cells[0].paragraphs[0].text, "Cell 1")
        self.assertEqual(table.rows[1].cells[1].paragraphs[0].text, "Cell 2")
        runs = table.rows[1].cells[1].paragraphs[0].runs
        self.assertEqual(runs[0].text, "Cell ")
        self.assertFalse(runs[0].bold)
        self.assertEqual(runs[1].text, "2")
        self.assertTrue(runs[1].bold)


if __name__ == "__main__":
    unittest.main()
